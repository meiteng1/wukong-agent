import os
from docx import Document

def main():
    path = os.path.abspath('厦门轨道交通桥梁支座检查报告_自动生成.docx')
    print('Doc:', path, os.path.exists(path))
    if not os.path.exists(path):
        return
    doc = Document(path)
    tables = doc.tables
    def is_5col_header(t):
        if not t.rows:
            return False
        hdr = [c.text.strip() for c in t.rows[0].cells[:5]]
        return hdr == ['桥墩','构件','部位','缺陷类型','现场照片']
    t5 = [t for t in tables if is_5col_header(t)]
    for idx, t in enumerate(t5):
        rows = len(t.rows)
        sample = [t.cell(1, i).text for i in range(5)] if rows > 1 else []
        flag = sample[1] if sample else ''
        tag = '3.1(#梁/#墩)' if ('#梁' in flag or '#墩' in flag) else '3.2(支座系统)'
        print(f'Table{idx} [{tag}] rows: {rows} sample: {sample}')

if __name__ == '__main__':
    main()
